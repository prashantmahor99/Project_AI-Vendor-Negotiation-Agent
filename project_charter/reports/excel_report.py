from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import os


def generate_excel_report(comparison):

    os.makedirs("output", exist_ok=True)

    wb = Workbook()

    ws = wb.active

    ws.title = "Price Comparison"

    # -----------------------------
    # Styles
    # -----------------------------

    header_fill = PatternFill(
        start_color="1F4E78",
        end_color="1F4E78",
        fill_type="solid"
    )

    l1_fill = PatternFill(
        start_color="92D050",
        end_color="92D050",
        fill_type="solid"
    )

    thin = Side(style="thin")

    border = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin
    )

    # -----------------------------
    # Heading
    # -----------------------------

    ws.merge_cells("A1:D1")

    cell = ws["A1"]

    cell.value = "PROCUREMENT PRICE COMPARISON"

    cell.font = Font(
        bold=True,
        color="FFFFFF",
        size=16
    )

    cell.fill = header_fill

    cell.alignment = Alignment(horizontal="center")

    # -----------------------------
    # Table Heading
    # -----------------------------

    headings = [

        "Rank",

        "Vendor Name",

        "Unit Price",

        "Difference"

    ]

    row = 3

    for col, value in enumerate(headings, start=1):

        c = ws.cell(
            row=row,
            column=col
        )

        c.value = value

        c.font = Font(bold=True)

        c.fill = header_fill

        c.alignment = Alignment(horizontal="center")

        c.border = border

    # -----------------------------
    # Data
    # -----------------------------

    lowest = comparison[0]["price"]

    row = 4

    for item in comparison:

        diff = item["price"] - lowest

        ws.cell(row=row, column=1).value = item["rank"]

        ws.cell(row=row, column=2).value = item["vendor"]

        ws.cell(row=row, column=3).value = item["price"]

        ws.cell(row=row, column=4).value = diff

        for c in range(1, 5):

            ws.cell(row=row, column=c).border = border

        if item["rank"] == "L1":

            for c in range(1, 5):

                ws.cell(
                    row=row,
                    column=c
                ).fill = l1_fill

        row += 1

    # -----------------------------
    # Lowest Vendor
    # -----------------------------

    row += 2

    ws.cell(row=row, column=1).value = "Lowest Vendor"

    ws.cell(row=row, column=2).value = comparison[0]["vendor"]

    row += 1

    ws.cell(row=row, column=1).value = "Lowest Price"

    ws.cell(row=row, column=2).value = comparison[0]["price"]

    # -----------------------------
    # Width
    # -----------------------------

    ws.column_dimensions["A"].width = 12

    ws.column_dimensions["B"].width = 35

    ws.column_dimensions["C"].width = 18

    ws.column_dimensions["D"].width = 18

    output = "output/Price_Comparison.xlsx"

    wb.save(output)

    return output